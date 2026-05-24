"""数据导出服务 - Excel / PDF."""
import io
from datetime import datetime

from app.models import GroupInfo, Task, TeamReport, UserProfile


class ExportService:
    """导出画像、分组、任务、报告."""

    def export_profiles(self, profiles: list) -> bytes:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Profiles"
        ws.append(["用户ID", "专业", "知识分", "技能分", "协作分", "偏好角色"])
        for p in profiles:
            ws.append([p.user_id, p.major, p.knowledge_score, p.skill_score, p.collab_score, p.pref_role])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def export_groups(self, groups: list) -> bytes:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Groups"
        ws.append(["组ID", "组名", "成员", "均衡分", "平均技能"])
        for g in groups:
            ws.append([g.id, g.group_name, g.member_ids, g.balance_score, g.avg_skill])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def export_tasks(self, tasks: list) -> bytes:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Tasks"
        ws.append(["任务", "负责人", "难度", "进度", "状态"])
        for t in tasks:
            ws.append([t.task_name, t.assignee_id, t.difficulty, t.progress, t.status])
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    def export_report_pdf(self, report: TeamReport, group: GroupInfo) -> bytes:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        c.drawString(50, 800, "组队超脑（TeamMind AI）- 团队协作质量报告")
        c.drawString(50, 780, f"团队: {group.group_name}")
        c.drawString(50, 760, f"完成率: {report.completion_rate}%")
        c.drawString(50, 740, f"质量分: {report.quality_score}")
        c.drawString(50, 720, f"风险: {report.risk_level}")
        c.drawString(50, 700, f"生成时间: {datetime.utcnow().isoformat()}")
        c.save()
        return buf.getvalue()
